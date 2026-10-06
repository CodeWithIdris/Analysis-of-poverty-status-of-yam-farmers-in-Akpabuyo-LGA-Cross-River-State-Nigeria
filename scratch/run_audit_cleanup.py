# -*- coding: utf-8 -*-
import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

def audit_final_doc(path):
    print(f'=== AUDITING {path} ===')
    doc = docx.Document(path)
    full_text = ' '.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    
    checks = [
        ('No "Off-farm income" in text/tables', 'Off-farm income' not in full_text),
        ('No "Yam / farm only" in text/tables', 'Yam / farm only' not in full_text and 'Yam/farm only' not in full_text),
        ('No "Local cultivars only" in text/tables', 'Local cultivars only' not in full_text),
        ('No "Traditional tools only" in text/tables', 'Traditional tools only' not in full_text),
        ('No "income diversification" in text', 'income diversification' not in full_text),
        ('No "dependency burden" in text', 'dependency burden' not in full_text),
        ('No "modern farm capital" in text', 'modern farm capital' not in full_text),
        ('Preserved "Yes (Other source of income)"', 'Yes (Other source of income)' in full_text),
        ('Preserved "No (No other source of income)"', 'No (No other source of income)' in full_text),
        ('Preserved "Yes (Uses improved yam varieties)"', 'Yes (Uses improved yam varieties)' in full_text),
        ('Preserved "No (Does not use improved yam varieties)"', 'No (Does not use improved yam varieties)' in full_text),
        ('Preserved "Yes (Uses modern farm tools/improved technologies)"', 'Yes (Uses modern farm tools/improved technologies)' in full_text),
        ('Preserved "No (Does not use modern farm tools/improved technologies)"', 'No (Does not use modern farm tools/improved technologies)' in full_text),
        ('Corrected household-size text present', 'This descriptive difference indicates that poor households had more household members' in full_text),
        ('Corrected other income text present', 'Regarding other sources of income, 85.11% of non-poor households reported' in full_text),
        ('Corrected farming experience text present', 'Poor farmers nevertheless reported greater mean yam-farming experience' in full_text),
        ('Corrected final conclusion present', 'Overall, the poverty-status profile shows that the poor group was descriptively characterized' in full_text),
        ('Preserved N = 60, Poor = 13, Non-poor = 47', '60' in full_text and '13' in full_text and '47' in full_text),
        ('Preserved Mean PCHE = 20,289.03, Line = 13,526.02', '20,289.03' in full_text and '13,526.02' in full_text),
        ('Preserved Poor Mean PCHE = 12,079.49, Non-Poor = 22,559.76', '12,079.49' in full_text and '22,559.76' in full_text),
        ('Preserved Poor Mean Total Exp = 101,000.00, Non-Poor = 119,941.49', '101,000.00' in full_text and '119,941.49' in full_text),
    ]
    
    all_pass = True
    for desc, res in checks:
        status = 'PASS' if res else 'FAIL'
        if not res:
            all_pass = False
        print(f'  [{status}] {desc}')
    return all_pass

res1 = audit_final_doc('Objective_2_Poverty_Status_Profile_Akpabuyo_FINAL.docx')
res2 = audit_final_doc('Chapter_4_Objective_2_Integration_FINAL.docx')

print(f'\nOverall Final Audit: doc1={res1}, doc2={res2}')

