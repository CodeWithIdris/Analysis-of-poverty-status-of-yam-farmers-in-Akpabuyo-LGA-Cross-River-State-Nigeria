import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

merged_path = os.path.join("Chapter_4_Analysis", "Chapter_4_Merged.docx")
backup_path = os.path.join("Chapter_4_Analysis", "Chapter_4_Merged_ORIGINAL.docx")

if not os.path.exists(backup_path) and os.path.exists(merged_path):
    shutil.copyfile(merged_path, backup_path)
    print(f"Backed up {merged_path} to {backup_path}")

doc_paths = [
    os.path.join("Chapter_4_Analysis", "Analysis_1_Socio_Economic_Characteristics", "Analysis_1_Chapter_4_Results.docx"),
    os.path.join("Chapter_4_Analysis", "Analysis_2_Household_Expenditure", "Analysis_2_Chapter_4_Results.docx"),
    os.path.join("Chapter_4_Analysis", "Analysis_3_Poverty_Status", "Analysis_3_Chapter_4_Results.docx"),
    os.path.join("Chapter_4_Analysis", "Analysis_4_Factors_Poverty", "Analysis_4_Chapter_4_Results.docx"),
    os.path.join("Chapter_4_Analysis", "Analysis_5_Challenges", "Analysis_5_Chapter_4_Results.docx"),
]

# We can create a merged document by combining elements from the 5 docs
merged_doc = Document()

# Set standard 1 inch margins
for section in merged_doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Chapter Four Title Header
p_ch = merged_doc.add_paragraph()
p_ch.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
p_ch.paragraph_format.space_before = Pt(0)
p_ch.paragraph_format.space_after = Pt(4)
r_ch = p_ch.add_run("CHAPTER FOUR")
r_ch.font.name = 'Calibri'
r_ch.font.size = Pt(18)
r_ch.font.bold = True
r_ch.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

p_sub = merged_doc.add_paragraph()
p_sub.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_after = Pt(18)
r_sub = p_sub.add_run("RESULTS AND DISCUSSION")
r_sub.font.name = 'Calibri'
r_sub.font.size = Pt(14)
r_sub.font.bold = True
r_sub.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

for doc_idx, p in enumerate(doc_paths):
    d = Document(p)
    # Copy body elements
    for element in d.element.body:
        tag_name = element.tag.split('}')[-1]
        if tag_name != 'sectPr':
            merged_doc.element.body.append(element)
    print(f"Appended doc {doc_idx+1}: {p}")

merged_doc.save(merged_path)
print(f"Successfully saved updated {merged_path}")
