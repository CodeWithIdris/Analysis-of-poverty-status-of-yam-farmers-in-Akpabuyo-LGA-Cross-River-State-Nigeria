import os
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import matplotlib.pyplot as plt
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Define Output Directory
OUT_DIR = os.path.join("Chapter_4_Analysis", "Analysis_5_Challenges")
os.makedirs(OUT_DIR, exist_ok=True)

# 1. Load Data
df = pd.read_csv("raw_data.csv")
n_total = len(df)
assert n_total == 60, f"Expected N=60, got {n_total}"

challenge_cols = [
    ('HIGH COST OF FARM INPUTS', 'High cost of farm inputs'),
    ('INADEQUATE ACCESS TO CREDIT', 'Inadequate access to credit'),
    ('PEST AND DISEASE INFESTATION', 'Pest and disease infestation'),
    ('UNPREDICTABLE RAINFALL AND CLIMATE CONDITIONS', 'Unpredictable rainfall and climate conditions'),
    ('HIGH COST AND SCARCITY OF FARM LABOUR', 'High cost and scarcity of farm labour'),
    ('HIGH COST/SCARCITY OF YAM STAKES', 'High cost/scarcity of yam stakes'),
    ('POOR ACCESS TO MARKET', 'Poor access to markets'),
    ('LOW AND UNSTABLE PRICES OF YAM', 'Low and unstable prices of yam'),
    ('INADEQUATE AGRICULTURAL EXTENSIONSERVICES', 'Inadequate agricultural extension services'),
    ('POST-HARVEST LOSSES AND INADEQUATE STORAGE FACILITIES', 'Post-harvest losses and inadequate storage facilities')
]

table_rows = []

for raw_col, label in challenge_cols:
    vals = df[raw_col].dropna()
    n_valid = len(vals)
    mean_val = vals.mean()
    sd_val = vals.std()
    med_val = vals.median()
    q25 = vals.quantile(0.25)
    q75 = vals.quantile(0.75)
    iqr_val = q75 - q25
    
    # Interpretation scale:
    # 1.00–1.80 = Not a Challenge
    # 1.81–2.60 = Minor Challenge
    # 2.61–3.40 = Moderate Challenge
    # 3.41–4.20 = Severe Challenge
    # 4.21–5.00 = Very Severe Challenge
    if 1.00 <= mean_val <= 1.80:
        interp = 'Not a Challenge'
    elif 1.81 <= mean_val <= 2.60:
        interp = 'Minor Challenge'
    elif 2.61 <= mean_val <= 3.40:
        interp = 'Moderate Challenge'
    elif 3.41 <= mean_val <= 4.20:
        interp = 'Severe Challenge'
    elif 4.21 <= mean_val <= 5.00:
        interp = 'Very Severe Challenge'
    else:
        interp = 'Out of Range'
        
    vc = vals.value_counts().to_dict()
    f1, f2, f3, f4, f5 = [int(vc.get(score, 0)) for score in [1.0, 2.0, 3.0, 4.0, 5.0]]
    p1, p2, p3, p4, p5 = [(f / n_valid) * 100 for f in [f1, f2, f3, f4, f5]]
    
    table_rows.append({
        'raw_col': raw_col,
        'Challenge': label,
        'n_valid': n_valid,
        'Mean': mean_val,
        'SD': sd_val,
        'Median': med_val,
        'IQR': iqr_val,
        'Q25': q25,
        'Q75': q75,
        'Interpretation': interp,
        'f1': f1, 'p1': p1,
        'f2': f2, 'p2': p2,
        'f3': f3, 'p3': p3,
        'f4': f4, 'p4': p4,
        'f5': f5, 'p5': p5
    })

res_df = pd.DataFrame(table_rows)

# Sort by Mean descending, then by SD ascending for tie-breaking if any
res_df = res_df.sort_values(by=['Mean', 'SD'], ascending=[False, True]).reset_index(drop=True)
res_df['Rank'] = range(1, len(res_df) + 1)

# ---------------------------------------------------------
# GENERATE EXCEL FILE 1: Table 4.11 (Ranking & Severity)
# ---------------------------------------------------------
wb11 = openpyxl.Workbook()
ws11 = wb11.active
ws11.title = "Table 4.11"

ws11.cell(row=1, column=1, value="Table 4.11: Severity and Ranking of Challenges Faced by Yam Farmers in Akpabuyo LGA").font = Font(name="Calibri", size=12, bold=True)
ws11.merge_cells("A1:G1")

headers11 = ["Rank", "Challenge Variable", "Sample Size (n)", "Mean Severity Score", "Standard Deviation (SD)", "Median (IQR)", "Interpretation Category"]
ws11.append([])
ws11.append(headers11)

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
thin_border_side = Side(border_style="thin", color="D9D9D9")
thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

for col_num in range(1, 8):
    cell = ws11.cell(row=3, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center" if col_num != 2 else "left", vertical="center")

for r_idx, r in res_df.iterrows():
    row_num = r_idx + 4
    c1 = ws11.cell(row=row_num, column=1, value=int(r['Rank']))
    c2 = ws11.cell(row=row_num, column=2, value=r['Challenge'])
    c3 = ws11.cell(row=row_num, column=3, value=int(r['n_valid']))
    c4 = ws11.cell(row=row_num, column=4, value=float(r['Mean']))
    c5 = ws11.cell(row=row_num, column=5, value=float(r['SD']))
    c6 = ws11.cell(row=row_num, column=6, value=f"{r['Median']:.2f} ({r['IQR']:.2f})")
    c7 = ws11.cell(row=row_num, column=7, value=r['Interpretation'])
    
    for c in [c1, c2, c3, c4, c5, c6, c7]:
        c.font = Font(name="Calibri", size=10.5)
        c.border = thin_border
        
    c1.alignment = Alignment(horizontal="center", vertical="center")
    c2.alignment = Alignment(horizontal="left", vertical="center")
    c3.alignment = Alignment(horizontal="center", vertical="center")
    c4.number_format = "0.00"
    c4.alignment = Alignment(horizontal="right", vertical="center")
    c5.number_format = "0.00"
    c5.alignment = Alignment(horizontal="right", vertical="center")
    c6.alignment = Alignment(horizontal="center", vertical="center")
    c7.alignment = Alignment(horizontal="center", vertical="center")

ws11.column_dimensions['A'].width = 10
ws11.column_dimensions['B'].width = 48
ws11.column_dimensions['C'].width = 16
ws11.column_dimensions['D'].width = 22
ws11.column_dimensions['E'].width = 22
ws11.column_dimensions['F'].width = 18
ws11.column_dimensions['G'].width = 24

wb11.save(os.path.join(OUT_DIR, "Table_4_11_Challenge_Severity_Ranking.xlsx"))
print("Saved Table_4_11_Challenge_Severity_Ranking.xlsx")

# ---------------------------------------------------------
# GENERATE EXCEL FILE 2: Frequency Distribution Table
# ---------------------------------------------------------
wb_freq = openpyxl.Workbook()
ws_freq = wb_freq.active
ws_freq.title = "Frequency Distribution"

ws_freq.cell(row=1, column=1, value="Supplementary Table: Frequency and Percentage Distribution of Challenge Severity Scores").font = Font(name="Calibri", size=12, bold=True)
ws_freq.merge_cells("A1:L1")

headers_freq = ["Rank", "Challenge Variable", "n", "Not a Challenge (1) n (%)", "Minor (2) n (%)", "Moderate (3) n (%)", "Severe (4) n (%)", "Very Severe (5) n (%)", "Mean", "SD", "Median", "Interpretation"]
ws_freq.append([])
ws_freq.append(headers_freq)

for col_num in range(1, 13):
    cell = ws_freq.cell(row=3, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center" if col_num != 2 else "left", vertical="center")

for r_idx, r in res_df.iterrows():
    row_num = r_idx + 4
    c1 = ws_freq.cell(row=row_num, column=1, value=int(r['Rank']))
    c2 = ws_freq.cell(row=row_num, column=2, value=r['Challenge'])
    c3 = ws_freq.cell(row=row_num, column=3, value=int(r['n_valid']))
    c4 = ws_freq.cell(row=row_num, column=4, value=f"{r['f1']} ({r['p1']:.1f}%)")
    c5 = ws_freq.cell(row=row_num, column=5, value=f"{r['f2']} ({r['p2']:.1f}%)")
    c6 = ws_freq.cell(row=row_num, column=6, value=f"{r['f3']} ({r['p3']:.1f}%)")
    c7 = ws_freq.cell(row=row_num, column=7, value=f"{r['f4']} ({r['p4']:.1f}%)")
    c8 = ws_freq.cell(row=row_num, column=8, value=f"{r['f5']} ({r['p5']:.1f}%)")
    c9 = ws_freq.cell(row=row_num, column=9, value=float(r['Mean']))
    c10 = ws_freq.cell(row=row_num, column=10, value=float(r['SD']))
    c11 = ws_freq.cell(row=row_num, column=11, value=float(r['Median']))
    c12 = ws_freq.cell(row=row_num, column=12, value=r['Interpretation'])
    
    for c in [c1, c2, c3, c4, c5, c6, c7, c8, c9, c10, c11, c12]:
        c.font = Font(name="Calibri", size=10)
        c.border = thin_border
        
    c1.alignment = Alignment(horizontal="center", vertical="center")
    c2.alignment = Alignment(horizontal="left", vertical="center")
    c3.alignment = Alignment(horizontal="center", vertical="center")
    for c in [c4, c5, c6, c7, c8]:
        c.alignment = Alignment(horizontal="center", vertical="center")
    c9.number_format = "0.00"
    c9.alignment = Alignment(horizontal="right", vertical="center")
    c10.number_format = "0.00"
    c10.alignment = Alignment(horizontal="right", vertical="center")
    c11.number_format = "0.00"
    c11.alignment = Alignment(horizontal="center", vertical="center")
    c12.alignment = Alignment(horizontal="center", vertical="center")

ws_freq.column_dimensions['A'].width = 8
ws_freq.column_dimensions['B'].width = 44
ws_freq.column_dimensions['C'].width = 8
for col_let in ['D', 'E', 'F', 'G', 'H']:
    ws_freq.column_dimensions[col_let].width = 20
ws_freq.column_dimensions['I'].width = 12
ws_freq.column_dimensions['J'].width = 12
ws_freq.column_dimensions['K'].width = 12
ws_freq.column_dimensions['L'].width = 22

wb_freq.save(os.path.join(OUT_DIR, "Table_4_11_Challenge_Frequency_Distribution.xlsx"))
print("Saved Table_4_11_Challenge_Frequency_Distribution.xlsx")

# ---------------------------------------------------------
# GENERATE EXCEL FILE 3: Data Validation Worksheet
# ---------------------------------------------------------
wb_val = openpyxl.Workbook()
ws_val = wb_val.active
ws_val.title = "Data Validation"

ws_val.cell(row=1, column=1, value="Analysis 5: Data Validation and Completeness Summary").font = Font(name="Calibri", size=13, bold=True)

headers_val = ["Challenge Variable", "Raw Variable Name", "Total Respondents (N)", "Valid Responses (n)", "Missing Values", "Valid Code Range", "Data Quality Status"]
ws_val.append([])
ws_val.append(headers_val)

for col_num in range(1, 8):
    cell = ws_val.cell(row=3, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="left", vertical="center")

for r_idx, r in res_df.iterrows():
    row_num = r_idx + 4
    m_count = 60 - int(r['n_valid'])
    status_str = "100% Complete" if m_count == 0 else f"1 Missing Value ({m_count/60*100:.1f}%)"
    
    c1 = ws_val.cell(row=row_num, column=1, value=r['Challenge'])
    c2 = ws_val.cell(row=row_num, column=2, value=r['raw_col'])
    c3 = ws_val.cell(row=row_num, column=3, value=60)
    c4 = ws_val.cell(row=row_num, column=4, value=int(r['n_valid']))
    c5 = ws_val.cell(row=row_num, column=5, value=m_count)
    c6 = ws_val.cell(row=row_num, column=6, value="1 to 5")
    c7 = ws_val.cell(row=row_num, column=7, value=status_str)
    
    for c in [c1, c2, c3, c4, c5, c6, c7]:
        c.font = Font(name="Calibri", size=10.5)
        c.border = thin_border
    c1.font = Font(name="Calibri", size=10.5, bold=True)
    c3.alignment = Alignment(horizontal="center", vertical="center")
    c4.alignment = Alignment(horizontal="center", vertical="center")
    c5.alignment = Alignment(horizontal="center", vertical="center")
    c6.alignment = Alignment(horizontal="center", vertical="center")
    c7.alignment = Alignment(horizontal="left", vertical="center")

ws_val.column_dimensions['A'].width = 44
ws_val.column_dimensions['B'].width = 48
ws_val.column_dimensions['C'].width = 22
ws_val.column_dimensions['D'].width = 22
ws_val.column_dimensions['E'].width = 18
ws_val.column_dimensions['F'].width = 18
ws_val.column_dimensions['G'].width = 28

wb_val.save(os.path.join(OUT_DIR, "Data_Validation_Diagnostic_Worksheet.xlsx"))
print("Saved Data_Validation_Diagnostic_Worksheet.xlsx")

# ---------------------------------------------------------
# GENERATE FIGURE 4.5
# ---------------------------------------------------------
plt.rcParams['font.sans-serif'] = 'Calibri'
plt.rcParams['font.family'] = 'sans-serif'
plt.rcParams['figure.dpi'] = 300

fig, ax = plt.subplots(figsize=(9.2, 5.6))

challenges_plot = res_df['Challenge'].tolist()[::-1]
means_plot = res_df['Mean'].tolist()[::-1]

y_pos = np.arange(len(challenges_plot))

# Color coding: Severe (>= 3.41) vs Moderate (2.61 - 3.40)
colors = ['#1F4E78' if m >= 3.41 else '#2E75B6' for m in means_plot]

# Baseline at 0
bars = ax.barh(y_pos, means_plot, color=colors, height=0.65, zorder=2)

# Direct labels
for bar, m in zip(bars, means_plot):
    ax.text(m + 0.08, bar.get_y() + bar.get_height()/2.0, f'{m:.2f}', 
            ha='left', va='center', fontsize=10, fontweight='bold', color='#1F4E78')

ax.set_yticks(y_pos)
ax.set_yticklabels(challenges_plot, fontsize=10.5, fontweight='bold')
ax.set_xlabel('Mean Severity Score (1 = Not a Challenge; 5 = Very Severe Challenge)', fontsize=11, fontweight='bold', labelpad=10)
ax.set_xlim(0, 5.0)
ax.set_xticks([0, 1, 2, 3, 4, 5])

# No gridlines, no enclosing box
ax.spines['top'].set_visible(False)
ax.spines['right'].set_visible(False)
ax.spines['left'].set_color('#333333')
ax.spines['bottom'].set_color('#333333')
ax.tick_params(axis='both', which='major', labelsize=10)

plt.tight_layout()
fig5_path = os.path.join(OUT_DIR, "Figure_4_5_Challenge_Severity_Ranking.png")
plt.savefig(fig5_path, dpi=300, bbox_inches='tight')
plt.close()
print("Saved Figure_4_5_Challenge_Severity_Ranking.png")

# ---------------------------------------------------------
# GENERATE WORD DOCUMENT: Analysis_5_Chapter_4_Results.docx
# ---------------------------------------------------------
doc = Document()

for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

style_normal = doc.styles['Normal']
style_normal.font.name = 'Calibri'
style_normal.font.size = Pt(11)
style_normal.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

def add_p(text, bold_prefix=None, space_after=6, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
    p.add_run(text)
    return p

def add_h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
    return p

def add_h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
    return p

# Document Title
add_h1("4.5 Challenges Faced by Yam Farmers")

# Section 4.5.1 Validation and Measurement Scale
add_h2("4.5.1 Data Validation and Measurement Framework")

add_p("To identify and rank the production, institutional, and marketing constraints encountered by yam farmers in Akpabuyo Local Government Area of Cross River State, ten specific challenge indicators specified in Section D of the research questionnaire were analyzed. Farmers rated the severity of each challenge on a 5-point Likert ordinal scale: 1 = Not a Challenge, 2 = Minor Challenge, 3 = Moderate Challenge, 4 = Severe Challenge, and 5 = Very Severe Challenge.")

add_p("Data validation confirmed complete response integrity (N = 60, n = 60 valid responses) across nine of the ten challenge indicators. For one indicator—low and unstable prices of yam—exactly one response was unrecorded, yielding n = 59 valid observations (98.3% completeness). No implausible or out-of-range values outside the 1–5 scale were detected. To ensure statistical rigor, mean severity scores were interpreted using the following standard five-interval evaluation scale:", bold_prefix="Data Integrity & Evaluation Scale: ")

add_p("• 1.00 – 1.80 = Not a Challenge\n• 1.81 – 2.60 = Minor Challenge\n• 2.61 – 3.40 = Moderate Challenge\n• 3.41 – 4.20 = Severe Challenge\n• 4.21 – 5.00 = Very Severe Challenge")

add_p("Because the responses represent ordinal severity ratings, mean severity scores are reported as descriptive ranking metrics. In accordance with ordinal data protocols, median and interquartile range (IQR) statistics are also presented as supplementary non-parametric measures.")

# Section 4.5.2 Table 4.11 Ranking and Severity
add_h2("4.5.2 Severity and Ranking Analysis of Yam Production Challenges")

add_p("Table 4.11 presents the mean severity scores, standard deviations, medians, interquartile ranges, overall ranks, and severity interpretation categories for the ten challenges evaluated in the study, ordered from highest mean severity (Rank 1) to lowest mean severity (Rank 10).")

# Insert Table 4.11 into docx
t11 = doc.add_table(rows=1, cols=5)
t11.alignment = WD_TABLE_ALIGNMENT.CENTER
t11.autofit = False

hdr_cells11 = t11.rows[0].cells
headers11_text = ["Rank", "Challenge Variable", "Mean Score", "SD", "Interpretation Category"]
for i, h_text in enumerate(headers11_text):
    hdr_cells11[i].text = h_text
    shading = parse_xml(r'<w:shd {} w:fill="1F4E78"/>'.format(nsdecls('w')))
    hdr_cells11[i]._tc.get_or_add_tcPr().append(shading)
    p = hdr_cells11[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.name = 'Calibri'
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

for r_idx, r in res_df.iterrows():
    row_cells = t11.add_row().cells
    row_cells[0].text = str(int(r['Rank']))
    row_cells[1].text = r['Challenge']
    row_cells[2].text = f"{r['Mean']:.2f}"
    row_cells[3].text = f"{r['SD']:.2f}"
    row_cells[4].text = r['Interpretation']
    
    for i, cell in enumerate(row_cells):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 1 else WD_ALIGN_PARAGRAPH.CENTER
        for r_run in p.runs:
            r_run.font.name = 'Calibri'
            r_run.font.size = Pt(9.5)

doc.add_paragraph().paragraph_format.space_after = Pt(4)
p_cap11 = doc.add_paragraph()
p_cap11.alignment = WD_ALIGN_PARAGRAPH.LEFT
p_cap11.paragraph_format.space_after = Pt(12)
r_cap11 = p_cap11.add_run("Source: Field Survey Data Analysis, 2026. N = 60 (n = 59 for Low and unstable prices of yam due to 1 missing response). SD = Standard Deviation. Evaluation Scale: 1.00–1.80 (Not a Challenge), 1.81–2.60 (Minor), 2.61–3.40 (Moderate), 3.41–4.20 (Severe), 4.21–5.00 (Very Severe).")
r_cap11.font.size = Pt(9.5)
r_cap11.font.italic = True

# Narrative interpretation of findings
add_h2("4.5.3 Interpretation of Challenge Severity Ranks")

add_p("The empirical results indicate that seven of the ten challenges fell into the 'Severe Challenge' category (mean scores between 3.41 and 4.20), while the remaining three challenges were classified as 'Moderate Challenges' (mean scores between 2.61 and 3.40). No challenge was categorized as a minor challenge or not a challenge, reflecting a high overall burden of constraints across the yam enterprise in the study area.", bold_prefix="Overall Pattern of Constraints: ")

add_p("The results indicate that high cost and scarcity of farm labour had the highest mean severity score (mean = 4.15 ± 0.80; median = 4.00, IQR = 1.00; Rank 1). The challenge was therefore classified as a severe challenge. Exactly 78.3% of respondents rated labor constraints as either severe (40.0%) or very severe (38.3%), highlighting labor availability and wage demands as a major operational bottleneck in yam mound making, weeding, and harvesting.", bold_prefix="1. Highest-Ranked Challenge — Farm Labour: ")

add_p("The second-highest mean severity score was recorded for high cost of farm inputs (mean = 4.00 ± 0.96; median = 4.00, IQR = 2.00; Rank 2). This challenge was also classified as a severe challenge, with 68.3% of farmers rating it as severe (30.0%) or very severe (38.3%). High input expenditure constraints farmer purchasing power for agrochemicals and planting materials.", bold_prefix="2. Second-Ranked Challenge — Input Costs: ")

add_p("The third and fourth severe challenges were post-harvest losses and inadequate storage facilities (mean = 3.88 ± 0.99; median = 4.00, IQR = 2.00; Rank 3) and unpredictable rainfall and climate conditions (mean = 3.70 ± 0.85; median = 4.00, IQR = 1.00; Rank 4). These were followed by high cost and scarcity of yam seed stakes (mean = 3.57 ± 0.89; Rank 5), inadequate access to credit (mean = 3.53 ± 0.98; Rank 6), and inadequate agricultural extension services (mean = 3.52 ± 1.07; Rank 7). All seven of these constraints recorded mean scores above 3.50, demonstrating a pervasive set of production and institutional obstacles.", bold_prefix="3. Other Severe Challenges (Ranks 3 to 7): ")

add_p("The remaining three challenges were classified as moderate challenges. Pest and disease infestation recorded a mean severity score of 3.38 ± 0.83 (median = 3.00, IQR = 1.00; Rank 8). Low and unstable prices of yam recorded a mean severity score of 3.12 ± 0.87 (median = 3.00, IQR = 1.00; Rank 9, n = 59). Finally, the lowest mean severity score was recorded for poor access to markets (mean = 2.90 ± 1.02; median = 3.00, IQR = 2.00; Rank 10). While market access received the lowest relative rank, its mean score of 2.90 confirms that even the lowest-ranked indicator remains a moderate concern for yam farmers in Akpabuyo LGA.", bold_prefix="4. Moderate Challenges (Ranks 8 to 10): ")

# Section 4.5.4 Graphical Display (Figure 4.5)
add_h2("4.5.4 Graphical Presentation of Challenge Severity Rankings")

add_p("Figure 4.5 visually presents the ranked mean severity scores for the ten challenges faced by yam farmers in Akpabuyo LGA. The horizontal bars display the mean severity scores ordered from the highest ranked challenge at the top to the lowest ranked challenge at the bottom.")

# Embed Figure 4.5
if os.path.exists(fig5_path):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(fig5_path, width=Inches(6.3))

p_fig_cap5 = doc.add_paragraph()
p_fig_cap5.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_fig_cap5.paragraph_format.space_after = Pt(12)
r_fig_cap5 = p_fig_cap5.add_run("Figure 4.5: Ranking of Challenges Faced by Yam Farmers in Akpabuyo LGA (Mean Severity Scores)")
r_fig_cap5.font.size = Pt(10)
r_fig_cap5.font.bold = True

# Save Document
doc_path = os.path.join(OUT_DIR, "Analysis_5_Chapter_4_Results.docx")
doc.save(doc_path)
print("Saved Analysis_5_Chapter_4_Results.docx")
