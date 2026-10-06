# -*- coding: utf-8 -*-
import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('../chapter 1-5.docx')

print("=== CHECKING PRELIMINARIES & ABSTRACT ===")
for i in range(0, 70):
    txt = doc.paragraphs[i].text.strip()
    if txt:
        print(f"P{i:03d}: {txt[:100]}")

print("\n=== CHECKING CHAPTER 1 OBJECTIVES & RESEARCH QUESTIONS ===")
for i in range(65, 110):
    txt = doc.paragraphs[i].text.strip()
    if any(k in txt for k in ['1.3', '1.4', '1.5', 'Research Question', 'Objective', 'Hypothesis', 'Yakurr', 'Akpabuyo']):
        print(f"P{i:03d}: {txt}")

print("\n=== CHECKING CHAPTER 3 DATA ANALYSIS METHODS ===")
for i in range(199, 331):
    txt = doc.paragraphs[i].text.strip()
    if any(k in txt for k in ['3.4', 'Method of Data Analysis', 'Analytical Technique', 'Model Specification', 'FGT', 'Logistic', 'Objective']):
        print(f"P{i:03d}: {txt}")

print("\n=== CHECKING CHAPTER 5 ===")
for i in range(467, 505):
    txt = doc.paragraphs[i].text.strip()
    if txt:
        print(f"P{i:03d}: {txt[:120]}")

