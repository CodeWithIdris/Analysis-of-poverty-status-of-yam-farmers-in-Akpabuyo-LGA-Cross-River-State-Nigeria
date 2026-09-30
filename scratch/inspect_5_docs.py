import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Let's inspect the 5 docs
doc_paths = [
    os.path.join("Chapter_4_Analysis", "Analysis_1_Socio_Economic_Characteristics", "Analysis_1_Chapter_4_Results.docx"),
    os.path.join("Chapter_4_Analysis", "Analysis_2_Household_Expenditure", "Analysis_2_Chapter_4_Results.docx"),
    os.path.join("Chapter_4_Analysis", "Analysis_3_Poverty_Status", "Analysis_3_Chapter_4_Results.docx"),
    os.path.join("Chapter_4_Analysis", "Analysis_4_Factors_Poverty", "Analysis_4_Chapter_4_Results.docx"),
    os.path.join("Chapter_4_Analysis", "Analysis_5_Challenges", "Analysis_5_Chapter_4_Results.docx"),
]

for p in doc_paths:
    d = Document(p)
    print(f"{p}: paragraphs={len(d.paragraphs)}, tables={len(d.tables)}")
