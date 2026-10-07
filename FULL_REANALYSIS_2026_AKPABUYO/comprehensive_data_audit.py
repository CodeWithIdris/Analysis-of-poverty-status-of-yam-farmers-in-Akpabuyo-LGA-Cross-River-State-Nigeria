import os
import sys
import docx
import pandas as pd
import numpy as np

# Set standard output encoding to utf-8
sys.stdout.reconfigure(encoding='utf-8')

print("================================================================================")
print("1. QUESTIONNAIRE TEXT EXTRACTION")
print("================================================================================")
doc_path = 'AKPABUYO_YAM_FARMERS_FINAL_PROJECT/05_QUESTIONNAIRE_AND_INSTRUMENTS/QUESTIONNAIRE_FINAL.docx'
doc = docx.Document(doc_path)
for p in doc.paragraphs:
    if p.text.strip():
        print(p.text)

for t_idx, table in enumerate(doc.tables):
    print(f"\n--- Questionnaire Table {t_idx+1} ---")
    for row in table.rows:
        row_vals = [c.text.strip() for c in row.cells]
        print(" | ".join(row_vals))

print("\n================================================================================")
print("2. RAW DATA COMPLETE PROFILE")
print("================================================================================")
df = pd.read_csv('raw_data.csv')
print(f"Dataset shape: {df.shape[0]} rows, {df.shape[1]} columns")

# Drop completely empty columns
valid_cols = [c for c in df.columns if not df[c].isnull().all()]
empty_cols = [c for c in df.columns if df[c].isnull().all()]
print(f"Empty/Section header columns ({len(empty_cols)}): {empty_cols}")
print(f"Valid data columns ({len(valid_cols)}): {valid_cols}")

print("\nDetailed breakdown of all valid data columns:")
for c in valid_cols:
    s = df[c]
    vc = s.value_counts(dropna=False).to_dict()
    print(f"\nColumn: [{c}]")
    print(f"  Type: {s.dtype} | Non-null: {s.count()}/60 | Null: {s.isnull().sum()}")
    print(f"  Min: {s.min()}, Max: {s.max()}, Mean: {s.mean():.4f}, Median: {s.median()}, SD: {s.std():.4f}")
    if len(vc) <= 15:
        print(f"  Value Counts: {vc}")
    else:
        print(f"  Distinct values ({len(vc)}): {sorted(list(vc.keys()))[:10]} ...")
